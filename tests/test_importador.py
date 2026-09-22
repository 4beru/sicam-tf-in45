"""Tests del importador Excel: plantilla, validación en seco y upsert."""
from __future__ import annotations

import sqlite3

import pytest
from openpyxl import load_workbook

from app.core import db, importador, plantilla


@pytest.fixture()
def conn(tmp_path):
    c = db.conectar(tmp_path / "test.db")
    yield c
    c.close()


@pytest.fixture()
def plantilla_ok(tmp_path):
    ruta = plantilla.generar_plantilla(tmp_path / "plantilla.xlsx")
    wb = load_workbook(ruta)

    ws = wb["maquinas"]
    ws.append(["PRUEBA-01", "Recta", "Juki DDL-8700", "Costura", "Alta", "2024-01-15", "2026-09-01"])

    ws = wb["produccion"]
    ws.append(["2026-09-13", "Costura", "Mañana", 500, 30])
    ws.append(["2026-09-13", "Costura", "Tarde", 450, 26])

    ws = wb["no_conformidades"]
    ws.append(["2026-09-13", "Tarde", "Costura", "Errores humanos",
               "Desfase de costura", "Rojas M.", "PRUEBA-01", 4])

    ws = wb["metas"]
    ws.append(["Tasa de prendas no conformes", 10.68, 4.95])

    wb.save(ruta)
    return ruta


def test_plantilla_tiene_hojas_y_validaciones(tmp_path):
    ruta = plantilla.generar_plantilla(tmp_path / "p.xlsx")
    wb = load_workbook(ruta)
    assert set(plantilla.HOJAS) <= set(wb.sheetnames)
    assert "Leeme" in wb.sheetnames
    # las hojas con catálogo llevan validación de datos
    assert len(wb["no_conformidades"].data_validations.dataValidation) > 0


def test_lectura_y_validacion_limpia(conn, plantilla_ok):
    datos = importador.leer(plantilla_ok)
    assert set(datos) == {"maquinas", "produccion", "no_conformidades", "metas"}
    val = importador.validar(datos, conn)
    assert val.ok, [e.detalle for e in val.errores]


def test_errores_de_validacion_no_tocan_bd(conn, plantilla_ok):
    wb = load_workbook(plantilla_ok)
    wb["no_conformidades"].append(["13/13/2026", "Tarde", "Costura",
                                   "Causa inventada", "X", "Nadie", "NO-EXISTE", 0])
    wb.save(plantilla_ok)

    datos = importador.leer(plantilla_ok)
    val = importador.validar(datos, conn)
    assert not val.ok
    detalles = [e.detalle for e in val.errores]
    assert any("fecha no válida" in d for d in detalles)
    assert any("causa" in d.lower() for d in detalles)
    assert any("maquina 'NO-EXISTE'" in d for d in detalles)
    assert any("cantidad" in d.lower() for d in detalles)
    # dry-run: la BD sigue vacía
    assert db.contar(conn)["maquinas"] == 0


def test_importar_y_reimportar_sin_duplicar(conn, plantilla_ok):
    datos = importador.leer(plantilla_ok)
    importador.importar(conn, datos)
    assert db.contar(conn)["maquinas"] == 1
    assert db.contar(conn)["produccion"] == 2

    # reimportar el mismo archivo: upsert, sin duplicados
    datos2 = importador.leer(plantilla_ok)
    res = importador.importar(conn, datos2)
    assert db.contar(conn)["maquinas"] == 1
    assert db.contar(conn)["produccion"] == 2
    assert res.resumen["maquinas"] == (0, 1)      # 0 nuevos, 1 actualizado
    assert res.resumen["produccion"] == (0, 2)

    # las NC sí se agregan (no tienen clave natural)
    assert db.contar(conn)["no_conformidades"] == 2


def test_archivo_sin_hojas_rechazado(tmp_path):
    from openpyxl import Workbook
    ruta = tmp_path / "vacio.xlsx"
    Workbook().save(ruta)
    with pytest.raises(ValueError):
        importador.leer(ruta)
