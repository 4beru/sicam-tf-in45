"""Generación de la plantilla Excel de carga de datos (plantilla_sicam.xlsx).

La plantilla es la puerta de entrada de la empresa al sistema: una hoja por
entidad, cabeceras fijas y listas desplegables (validación de datos) para que
sea imposible escribir categorías con otra ortografía.
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import db

# hoja -> (columnas, columnas con lista desplegable -> catálogo)
HOJAS: dict[str, tuple[list[str], dict[int, list[str]]]] = {
    "maquinas": (
        ["codigo", "tipo", "marca_modelo", "area", "criticidad", "fecha_compra", "ultima_calibracion"],
        {4: db.AREAS, 5: db.CRITICIDAD},
    ),
    "operarios": (
        ["codigo", "nombre", "area", "proceso"],
        {3: db.AREAS, 4: db.PROCESOS},
    ),
    "produccion": (
        ["fecha", "proceso", "turno", "prendas_inspeccionadas", "nc_detectadas"],
        {2: db.PROCESOS, 3: db.TURNOS},
    ),
    "no_conformidades": (
        ["fecha", "turno", "proceso", "causa", "defecto", "operario", "maquina", "cantidad"],
        {2: db.TURNOS, 3: db.PROCESOS, 4: db.CAUSAS},
    ),
    "frecuencias": (
        ["maquina", "actividad", "frecuencia_dias", "responsable"],
        {},
    ),
    "metas": (
        ["indicador", "as_is", "to_be"],
        {},
    ),
}

TEAL = "0D9488"
_MESES = ["", "enero", "febrero", "marzo", "abril", "mayo", "junio",
          "julio", "agosto", "setiembre", "octubre", "noviembre", "diciembre"]


def generar_plantilla(ruta: Path | str) -> Path:
    """Escribe la plantilla con instrucciones, cabeceras y validaciones."""
    from openpyxl.worksheet.datavalidation import DataValidation

    ruta = Path(ruta)
    wb = Workbook()

    # ---- hoja de instrucciones
    info = wb.active
    info.title = "Leeme"
    lineas = [
        ("SICAM — Plantilla de carga de datos", True),
        ("", False),
        ("Cómo usar esta plantilla:", True),
        ("1. Cada hoja es una tabla del sistema: no cambies ni muevas los títulos de la primera fila.", False),
        ("2. Las columnas con lista desplegable solo aceptan los valores del desplegable.", False),
        ("3. Fechas en formato AAAA-MM-DD (ej. 2026-09-13) o DD/MM/AAAA.", False),
        ("4. Reimportar un archivo no duplica datos: las filas se actualizan por su clave", False),
        ("   (código de máquina/operario, fecha+proceso+turno, indicador…).", False),
        ("5. El sistema valida todo antes de importar: si hay errores te dirá fila y columna", False),
        ("   exactas, y no tocará la base de datos hasta que todo esté correcto.", False),
        ("", False),
        ("Puedes llenar solo las hojas que necesites y borrar las que no uses.", False),
    ]
    for i, (texto, negrita) in enumerate(lineas, start=1):
        c = info.cell(row=i, column=1, value=texto)
        c.font = Font(bold=negrita, size=14 if i == 1 else 10.5,
                      color="FFFFFF" if i == 1 else "1E293B")
    info.column_dimensions["A"].width = 105

    fino = Side(style="thin", color="CBD5E1")
    for hoja, (columnas, desplegables) in HOJAS.items():
        ws = wb.create_sheet(hoja)
        ws.freeze_panes = "A2"
        for j, nombre in enumerate(columnas, start=1):
            c = ws.cell(row=1, column=j, value=nombre)
            c.font = Font(bold=True, color="FFFFFF", size=10.5)
            c.fill = PatternFill("solid", fgColor=TEAL)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(bottom=fino)
            ws.column_dimensions[get_column_letter(j)].width = max(14, len(nombre) + 6)
        # listas desplegables (aplicadas a 2000 filas)
        for indice_col, valores in desplegables.items():
            col = get_column_letter(indice_col)
            dv = DataValidation(
                type="list",
                formula1='"' + ",".join(v.replace(",", " ") for v in valores) + '"',
                allow_blank=True,
                showErrorMessage=True,
                errorTitle="Valor no válido",
                error="Elige un valor de la lista desplegable.",
            )
            ws.add_data_validation(dv)
            dv.add(f"{col}2:{col}2001")

    wb.save(ruta)
    return ruta


def exportar_no_conformidades(conn, ruta: Path | str, desde: str | None = None) -> int:
    """Exporta el registro de NC a XLSX con las mismas columnas de la plantilla."""
    from openpyxl import Workbook

    columnas = HOJAS["no_conformidades"][0]
    sql = ("SELECT fecha, turno, proceso, causa, defecto, operario, maquina, cantidad "
           "FROM no_conformidades")
    params: list = []
    if desde:
        sql += " WHERE fecha >= ?"
        params.append(desde)
    sql += " ORDER BY fecha DESC, id DESC"
    filas = conn.execute(sql, params).fetchall()

    ruta = Path(ruta)
    wb = Workbook()
    ws = wb.active
    ws.title = "no_conformidades"
    ws.freeze_panes = "A2"
    for j, nombre in enumerate(columnas, start=1):
        c = ws.cell(row=1, column=j, value=nombre)
        c.font = Font(bold=True, color="FFFFFF", size=10.5)
        c.fill = PatternFill("solid", fgColor=TEAL)
        ws.column_dimensions[get_column_letter(j)].width = max(14, len(nombre) + 6)
    for i, fila in enumerate(filas, start=2):
        for j, valor in enumerate(tuple(fila), start=1):
            ws.cell(row=i, column=j, value=valor)
    wb.save(ruta)
    return len(filas)
